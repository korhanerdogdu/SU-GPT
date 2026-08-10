from __future__ import annotations

import hashlib
import json
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from modules.syllabus_pipeline import (
    detect_attachment_format,
    extract_docx,
    make_rag_chunks,
    parse_legacy_html,
    parse_new_app_html,
)
from modules.retrieval_policy import build_metadata_filter
from modules.catalog_retriever import _candidate_filters, _extract_course_ids
from retrieval_lab.corpus import load_corpus
from scripts.validate_syllabus_corpus import validate_term
from scripts import scrape_syllabi as crawler
from scripts.scrape_course_schedule import parse_course_header


def test_schedule_header_accepts_multi_letter_course_suffix():
    parsed = parse_course_header(
        "Readings and Research in Political Science: Feminist Political Theory - 23527 - POLS 700FT - 0"
    )
    assert parsed.course_id == "POLS 700FT"
    assert parsed.course_number == "700FT"
    assert parsed.component == "Primary"


def test_parse_legacy_page_with_attachment_and_instructor():
    html = b"""
    <html><main>
      <h1>AL 102, Academic Literacies, 202502</h1>
      <p><strong>Instructor:</strong> Naime Meltem Aygunes,
        <a href="mailto:meltemb@sabanciuniv.edu">meltemb@sabanciuniv.edu</a></p>
      <h2>Syllabus</h2><p>Students develop academic reading and writing skills.</p>
      <h2>Attachments</h2>
      <a href="https://sucourse.sabanciuniv.edu/plus/syllabusdownload.php?context=1&amp;filename=AL102.docx">
        AL102.docx
      </a>
    </main></html>
    """
    parsed = parse_legacy_html(
        html,
        base_url="https://www.sabanciuniv.edu/syllabus/?crn=10001&term=202502",
        expected_term="202502",
        expected_course_id="AL 102",
        expected_section="A",
    )
    assert parsed["page_state"] == "published"
    assert parsed["identity_verified"] is True
    assert parsed["instructor_emails"] == ["meltemb@sabanciuniv.edu"]
    assert parsed["attachments"][0]["filename"] == "AL102.docx"
    assert "academic reading" in parsed["inline_text"]


def test_parse_new_app_structured_page():
    html = b"""
    <html><main>
      <h1>CS 301 - Algorithms</h1>
      <table>
        <tr><th>Semester</th><td>Spring 2025-2026</td></tr>
        <tr><th>Section</th><td>A</td></tr>
        <tr><th>Instructor</th><td>Example Instructor</td></tr>
        <tr><th>Course Objective</th><td>Analyze algorithm correctness and complexity.</td></tr>
      </table>
      <h2>Learning Outcomes</h2><p>Students compare graph algorithms and prove bounds.</p>
    </main></html>
    """
    parsed = parse_new_app_html(
        html,
        base_url="https://apps.sabanciuniv.edu/courses/syllabus/view.php?cn=301&sc=CS&term=202502",
        expected_term="202502",
        expected_course_id="CS 301",
        expected_section="A",
    )
    assert parsed["page_state"] == "published"
    assert parsed["identity_verified"] is True
    assert "Example Instructor" in parsed["instructor_names"]
    assert parsed["fields"]["course objective"].startswith("Analyze")


def test_identity_verified_empty_legacy_page_is_only_a_candidate():
    html = b"<html><main><h1>CS 204, Advanced Programming, 202502</h1><h2>Syllabus</h2><h2>Attachments</h2></main></html>"
    parsed = parse_legacy_html(
        html,
        base_url="https://www.sabanciuniv.edu/syllabus/?crn=10002&term=202502",
        expected_term="202502",
        expected_course_id="CS 204",
        expected_section="A",
    )
    assert parsed["page_state"] == "not_published_candidate"
    assert parsed["identity_verified"] is True


def test_future_missing_page_uses_verified_endpoint_without_accepting_wrong_term_title():
    html = b"""
    <html><head><title>Syllabus-AL102-A2-202503</title></head>
    <main><h4>Syllabus Application</h4><p>Syllabus not found for this course section.</p></main></html>
    """
    parsed = parse_new_app_html(
        html,
        base_url="https://apps.sabanciuniv.edu/courses/syllabus/view.php?cn=102&sc=AL&section=A2&term=202601&view=public",
        expected_term="202601",
        expected_course_id="AL 102",
        expected_section="A2",
        expected_crn="10267",
    )
    assert parsed["page_state"] == "not_published_candidate"
    assert parsed["identity_verified"] is False
    assert parsed["identity_term_match"] is False
    assert parsed["endpoint_target_verified"] is True


def test_primary_section_ignores_prose_that_begins_with_section_word():
    html = b"""
    <html><head><title>Syllabus-CS412-202502</title></head><main>
      <h3>CS 412 Machine Learning</h3><p>Semester Spring 2025-2026</p>
      <h5>Learning Outcomes</h5><p>This section introduces robust model evaluation and optimization.</p>
    </main></html>
    """
    parsed = parse_new_app_html(
        html,
        base_url="https://apps.sabanciuniv.edu/courses/syllabus/view.php?cn=412&sc=CS&section=0&term=202502&view=public",
        expected_term="202502",
        expected_course_id="CS 412",
        expected_section="0",
    )
    assert parsed["page_state"] == "published"
    assert parsed["identity_section_match"] is None


def test_non_published_record_never_generates_rag_chunks():
    record = {
        "term": "202601",
        "crn": "10267",
        "course_id": "AL 102",
        "publication_state": "not_published_confirmed",
        "sources": [{"inline_text": "stale syllabus content from another term"}],
        "attachments": [],
    }
    assert make_rag_chunks(record, {}) == []


def test_crawler_requires_two_verified_empty_attempts_on_both_surfaces():
    original = crawler.fetch_page

    def fake_fetch(_session, *, source_system, url, **_kwargs):
        return {
            "source_system": source_system,
            "requested_url": url,
            "final_url": url,
            "http_status": 200,
            "fetched_at": "2026-08-10T00:00:00Z",
            "page_sha256": ("a" if source_system == "legacy" else "b") * 64,
            "page_state": "not_published_candidate",
            "identity_verified": True,
            "attachments": [],
            "inline_text": "",
            "instructor_names": [],
            "instructor_emails": [],
        }

    crawler.fetch_page = fake_fetch
    try:
        with tempfile.TemporaryDirectory() as directory:
            record, _parses, chunks = crawler.crawl_one(
                {
                    "term": "202601",
                    "crn": "10001",
                    "course_id": "CS 301",
                    "title": "Algorithms",
                    "subject": "CS",
                    "course_number": "301",
                    "section": "A",
                    "component": "Primary",
                },
                output_root=Path(directory),
                timeout=1,
                delay=0,
                confirm_empty=True,
                max_attachment_bytes=1024,
            )
    finally:
        crawler.fetch_page = original
    assert record["publication_state"] == "not_published_confirmed"
    assert record["capture_state"] == "not_applicable"
    assert all(source["empty_confirmation_agrees"] for source in record["sources"])
    assert chunks == []


def test_magic_detection_rejects_html_disguised_as_pdf():
    detected = detect_attachment_format(b"<!doctype html><html>Login</html>", "syllabus.pdf", "application/pdf")
    assert detected["valid"] is False
    assert detected["kind"] == "html"


def test_docx_extraction_keeps_tables():
    from docx import Document

    with tempfile.TemporaryDirectory() as directory:
        path = Path(directory) / "syllabus.docx"
        document = Document()
        document.add_heading("Assessment", level=1)
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Item"
        table.cell(0, 1).text = "Weight"
        table.cell(1, 0).text = "Midterm"
        table.cell(1, 1).text = "40%"
        document.save(path)
        parsed = extract_docx(path)
    assert parsed["parse_state"] == "parsed"
    text = "\n".join(unit["text"] for unit in parsed["units"])
    assert "Midterm" in text
    assert "40%" in text


def test_rag_chunks_have_stable_parent_metadata():
    record = {
        "term": "202502",
        "crn": "10001",
        "course_id": "AL 102",
        "course_title": "Academic Literacies",
        "subject": "AL",
        "course_number": "102",
        "section": "A",
        "component": "Primary",
        "instructor_names": ["Naime Meltem Aygunes"],
        "instructor_emails": ["meltemb@sabanciuniv.edu"],
        "publication_state": "published",
        "sources": [
            {
                "source_system": "legacy",
                "requested_url": "https://www.sabanciuniv.edu/syllabus/?crn=10001&term=202502",
                "raw_html_path": "raw/202502/10001/legacy.html",
                "inline_text": "Learning outcomes and grading policy are described here in sufficient detail for retrieval.",
            }
        ],
        "attachments": [],
    }
    first = make_rag_chunks(record, {})
    second = make_rag_chunks(record, {})
    assert first == second
    assert first[0]["data_role"] == "course_syllabus"
    assert first[0]["term"] == "202502"
    assert first[0]["crn"] == "10001"


def test_syllabus_policy_scopes_to_the_syllabus_role():
    assert build_metadata_filter("syllabus", {}) == {"data_role": "course_syllabus"}


def test_syllabus_filter_keeps_course_code_without_false_program_scope():
    filters = _candidate_filters("ECON 301 syllabus policies 202502")
    assert filters
    assert all("program" not in json.dumps(item) for item in filters)
    assert any("ECON 301" in json.dumps(item) for item in filters)
    assert _extract_course_ids("POLS 700FT and MAT 68004 syllabi") == ["POLS 700FT", "MAT 68004"]


def test_default_disk_corpus_loads_only_syllabus_chunks_not_ledgers():
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        syllabi = root / "syllabi"
        _write_jsonl(
            syllabi / "202502.jsonl",
            [{"text": "accounting ledger row", "chunk_id": "ledger-should-not-load"}],
        )
        _write_jsonl(
            syllabi / "202502.chunks.jsonl",
            [
                {
                    "text": "Official CS 301 syllabus learning outcomes.",
                    "chunk_id": "course_syllabus:202502:12345:0:0:abc",
                    "data_role": "course_syllabus",
                    "term": "202502",
                    "course_id": "CS 301",
                }
            ],
        )
        corpus = load_corpus(root)
    assert len(corpus.chunks) == 1
    assert corpus.chunks[0].chunk_id.startswith("course_syllabus:")
    assert corpus.chunks[0].curriculum_term == "202502"


def _write_jsonl(path: Path, rows: list[dict]) -> str:
    payload = "".join(json.dumps(row, ensure_ascii=False, sort_keys=True) + "\n" for row in rows).encode()
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(payload)
    return hashlib.sha256(payload).hexdigest()


def test_validator_rejects_self_attested_fixture_without_raw_evidence():
    with tempfile.TemporaryDirectory() as directory:
        root = Path(directory)
        schedule_root = root / "schedule"
        data_root = root / "syllabi"
        schedule = [
            {
                "term": "202502",
                "course_id": "CS 301",
                "title": "Algorithms",
                "subject": "CS",
                "course_number": "301",
                "section": "A",
                "component": "Primary",
                "crn": "12345",
            }
        ]
        ledger = [
            {
                "term": "202502",
                "term_code": "202502",
                "course_id": "CS 301",
                "course_title": "Algorithms",
                "subject": "CS",
                "course_number": "301",
                "section": "A",
                "component": "Primary",
                "crn": "12345",
                "instructor_names": ["Example Instructor"],
                "instructor_emails": ["example@sabanciuniv.edu"],
                "publication_state": "published",
                "capture_state": "complete",
                "status_reason": "published content found",
                "sources": [
                    {
                        "page_state": "published",
                        "http_status": 200,
                        "fetched_at": "2026-08-10T00:00:00Z",
                        "final_url": "https://apps.sabanciuniv.edu/courses/syllabus/view.php?cn=301&sc=CS&term=202502",
                        "page_sha256": "a" * 64,
                        "attachments": [],
                    }
                ],
                "attachments": [],
                "source_authority": "Sabanci University",
                "scraped_at": "2026-08-10T00:00:00Z",
                "syllabus_id": "syllabus:202502:12345",
                "chunk_id": "course_syllabus:202502:12345:overview",
                "text": "Official syllabus record for CS 301 Algorithms.",
            }
        ]
        chunks = [
            {
                "term": "202502",
                "course_id": "CS 301",
                "course_title": "Algorithms",
                "crn": "12345",
                "section": "A",
                "instructors": "Example Instructor",
                "locator": "Learning Outcomes",
                "source_url": "https://apps.sabanciuniv.edu/courses/syllabus/view.php?cn=301&sc=CS&term=202502",
                "data_role": "course_syllabus",
                "text": "Official syllabus for CS 301 Algorithms term 202502 CRN 12345. Learning outcomes cover graph algorithms and complexity analysis.",
                "chunk_id": "course_syllabus:202502:12345:0:0:abc",
            }
        ]
        _write_jsonl(schedule_root / "202502.jsonl", schedule)
        ledger_sha = _write_jsonl(data_root / "202502.jsonl", ledger)
        chunks_sha = _write_jsonl(data_root / "202502.chunks.jsonl", chunks)
        (data_root / "202502.manifest.json").write_text(
            json.dumps(
                {
                    "crawler_version": "test",
                    "schedule_crn_count": 1,
                    "ledger_record_count": 1,
                    "rag_chunk_count": 1,
                    "ledger_sha256": ledger_sha,
                    "chunks_sha256": chunks_sha,
                }
            ),
            encoding="utf-8",
        )
        result = validate_term("202502", data_root, schedule_root)
    assert result["score"] < 90.0
    assert result["passed"] is False
    assert any("raw HTML evidence" in blocker for blocker in result["blockers"])
    assert any("legacy+new_app evidence" in blocker for blocker in result["blockers"])
