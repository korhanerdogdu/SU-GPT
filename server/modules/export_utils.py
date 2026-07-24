from __future__ import annotations

"""CSV/XLSX generation for course history and deterministic audit tables."""

import csv
import io
from typing import Any, Iterable


def audit_rows(audit: dict[str, Any]) -> list[dict[str, Any]]:
    rows = []
    for item in audit.get("categories") or []:
        rows.append(
            {
                "type": "SU",
                "category": item.get("category"),
                "completed": item.get("completed_su_credits"),
                "required": item.get("required_su_credits"),
                "remaining": item.get("remaining_su_credits"),
                "courses": ", ".join(item.get("completed_courses") or []),
            }
        )
    for item in audit.get("ects_requirements") or []:
        rows.append(
            {
                "type": "ECTS",
                "category": item.get("category"),
                "completed": item.get("completed_ects"),
                "required": item.get("required_ects"),
                "remaining": item.get("remaining_ects"),
                "courses": "",
            }
        )
    return rows


def course_rows(courses: Iterable[dict[str, Any]]) -> list[dict[str, Any]]:
    rows = [
        {
            "code": course.get("code"),
            "title": course.get("title"),
            "status": course.get("status", "completed"),
            "su_credits": course.get("su_credits"),
            "ects": course.get("ects"),
            "engineering_ects": course.get("engineering_ects"),
            "basic_science_ects": course.get("basic_science_ects"),
        }
        for course in courses
    ]
    eligible = {"completed", "transfer", "exempted"}
    eligible_rows = [row for row in rows if str(row.get("status") or "completed").lower() in eligible]
    totals = {
        "su_credits": sum(float(row.get("su_credits") or 0) for row in eligible_rows),
        "ects": sum(float(row.get("ects") or 0) for row in eligible_rows),
        "engineering_ects": sum(float(row.get("engineering_ects") or 0) for row in eligible_rows),
        "basic_science_ects": sum(float(row.get("basic_science_ects") or 0) for row in eligible_rows),
    }
    for row in rows:
        row["total_su_credits"] = totals["su_credits"]
        row["total_ects"] = totals["ects"]
    rows.append(
        {
            "code": "TOTAL",
            "title": f"{len(eligible_rows)} credit-eligible courses",
            "status": "",
            **totals,
            "total_su_credits": totals["su_credits"],
            "total_ects": totals["ects"],
        }
    )
    return rows


def rows_to_csv(rows: list[dict[str, Any]]) -> bytes:
    if not rows:
        return b""
    output = io.StringIO(newline="")
    writer = csv.DictWriter(output, fieldnames=list(rows[0]))
    writer.writeheader()
    writer.writerows(rows)
    return output.getvalue().encode("utf-8-sig")


def rows_to_xlsx(rows: list[dict[str, Any]], *, sheet_name: str = "adviSU") -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as exc:
        raise RuntimeError("XLSX export requires openpyxl. Install server requirements first.") from exc

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name[:31]
    if rows:
        headers = list(rows[0])
        worksheet.append(headers)
        for row in rows:
            worksheet.append([row.get(header) for header in headers])
        header_fill = PatternFill("solid", fgColor="004B93")
        for cell in worksheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for column in worksheet.columns:
            width = min(max(len(str(cell.value or "")) for cell in column) + 2, 60)
            worksheet.column_dimensions[column[0].column_letter].width = width

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
